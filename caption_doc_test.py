from docx import Document
from docx.shared import Inches

# create a new document
doc = Document()

# add a table
table = doc.add_table(rows=3, cols=3)
cell = table.cell(0, 0)
cell.text = 'Row 0, Column 0'

# add a caption to the table
caption = table.add_caption('This is a caption')
# AttributeError: 'Table' object has no attribute 'add_caption'
# save the document
doc.save('my_table.docx')
